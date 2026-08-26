import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from extrator import (
    duplicate_previous_season_workbook,
    previous_season_token,
    write_seasons_manifest,
    workbooks_have_same_data,
)
from preditor import season_context


class SeasonRolloverTests(unittest.TestCase):
    def _create_workbook(
        self,
        path: Path,
        *,
        score: int = 2,
        sheet_name: str = "FUTSAL MASCULINO",
        styled: bool = False,
    ) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.append(["Equipa 1", "Golos 1", "Golos 2", "Equipa 2"])
        sheet.append(["Engenharia", score, 1, "Direito"])

        if styled:
            sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
            sheet["Z100"].fill = PatternFill(fill_type="solid", fgColor="00FF00")

        workbook.save(path)
        workbook.close()

    def test_previous_season_token(self):
        self.assertEqual(previous_season_token("26_27"), "25_26")
        self.assertEqual(previous_season_token("00_01"), "99_00")
        self.assertIsNone(previous_season_token("invalida"))

    def test_same_cell_data_ignores_formatting_and_trailing_empty_cells(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            first = Path(temp_dir) / "25_26.xlsx"
            second = Path(temp_dir) / "26_27.xlsx"
            self._create_workbook(first)
            self._create_workbook(second, styled=True)

            self.assertTrue(workbooks_have_same_data(first, second))

    def test_different_cell_data_is_not_equal(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            first = Path(temp_dir) / "25_26.xlsx"
            second = Path(temp_dir) / "26_27.xlsx"
            self._create_workbook(first, score=2)
            self._create_workbook(second, score=3)

            self.assertFalse(workbooks_have_same_data(first, second))

    def test_different_sheet_names_are_not_equal(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            first = Path(temp_dir) / "25_26.xlsx"
            second = Path(temp_dir) / "26_27.xlsx"
            self._create_workbook(first, sheet_name="FUTSAL MASCULINO")
            self._create_workbook(second, sheet_name="FUTSAL FEMININO")

            self.assertFalse(workbooks_have_same_data(first, second))

    def test_duplicate_rollover_is_blocked_until_a_cell_changes(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir)
            previous = data_dir / "Resultados Taça UA 25_26.xlsx"
            downloaded = data_dir / "export.xlsx"
            self._create_workbook(previous, score=2)
            self._create_workbook(downloaded, score=2, styled=True)

            self.assertEqual(
                duplicate_previous_season_workbook(
                    downloaded, data_dir, "26_27"
                ),
                previous,
            )

            self._create_workbook(downloaded, score=3)
            self.assertIsNone(
                duplicate_previous_season_workbook(
                    downloaded, data_dir, "26_27"
                )
            )

    def test_manifest_only_publishes_changed_seasons(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            data_dir = root / "data"
            csv_dir = root / "csv"
            manifest = root / "seasons.json"
            data_dir.mkdir()
            csv_dir.mkdir()

            previous = data_dir / "Resultados Taça UA 25_26.xlsx"
            current = data_dir / "Resultados Taça UA 26_27.xlsx"
            self._create_workbook(previous, score=2)
            self._create_workbook(current, score=2, styled=True)
            (csv_dir / "FUTSAL MASCULINO_25_26.csv").write_text(
                "Equipa 1,Equipa 2\n", encoding="utf-8"
            )
            (csv_dir / "FUTSAL MASCULINO_26_27.csv").write_text(
                "Equipa 1,Equipa 2\n", encoding="utf-8"
            )

            self.assertEqual(
                write_seasons_manifest(csv_dir, data_dir, manifest), ["25_26"]
            )
            self.assertEqual(
                json.loads(manifest.read_text(encoding="utf-8")),
                {"seasons": ["25_26"]},
            )

            self._create_workbook(current, score=3)
            self.assertEqual(
                write_seasons_manifest(csv_dir, data_dir, manifest),
                ["26_27", "25_26"],
            )

    def test_predictor_uses_august_season_rollover(self):
        self.assertEqual(
            season_context(datetime(2026, 7, 31)), ("25_26", "24_25", 2026)
        )
        self.assertEqual(
            season_context(datetime(2026, 8, 1)), ("26_27", "25_26", 2027)
        )


if __name__ == "__main__":
    unittest.main()
